from pathlib import Path
import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SRC = Path(os.environ.get("REPORT_SRC", ROOT / "output" / "video_diffusion_load_report.md"))
OUT = Path(os.environ.get("REPORT_OUT", ROOT / "output" / "Efficient_Video_Diffusion算法与负载分析_论文详述版.docx"))
PAPER_MODE = os.environ.get("PAPER_MODE", "0") == "1"

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT = "EAF1F8"
PALE = "F4F6F9"
GRAY = "5B6573"
GOLD = "A06A00"

def set_font(run, size=None, bold=None, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None: tcPr.append(shd)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m, v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        el = tcMar.find(qn(f"w:{m}")) or OxmlElement(f"w:{m}")
        el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa")
        if el.getparent() is None: tcMar.append(el)

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), BLUE); rPr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    r.append(rPr); t = OxmlElement("w:t"); t.text = text; r.append(t); hyperlink.append(r)
    paragraph._p.append(hyperlink)

def add_inline(p, text, base_size=10.5):
    url_re = re.compile(r"https?://[^\s)]+")
    pos = 0
    for m in url_re.finditer(text):
        if m.start() > pos:
            add_emphasis_runs(p, text[pos:m.start()], base_size)
        add_hyperlink(p, m.group(), m.group())
        pos = m.end()
    if pos < len(text): add_emphasis_runs(p, text[pos:], base_size)

def add_emphasis_runs(p, text, size):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part: continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        r = p.add_run(content); set_font(r, size=size, bold=bold, color=NAVY if bold else None)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); tblHeader = OxmlElement("w:tblHeader"); tblHeader.set(qn("w:val"), "true"); trPr.append(tblHeader)

def style_formula_paragraph(p):
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "EDF4FB"); pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left"); left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18"); left.set(qn("w:color"), BLUE); left.set(qn("w:space"), "8")
    pBdr.append(left); pPr.append(pBdr)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 "); set_font(run, 8.5, color=GRAY)
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve"); instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.extend([fldChar1, instrText, fldChar2])
    r2 = paragraph.add_run(" 页"); set_font(r2, 8.5, color=GRAY)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.82)
sec.left_margin = sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.35); sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.18
for name, size, before, after, color in [("Heading 1",16,14,7,NAVY),("Heading 2",13,11,5,BLUE),("Heading 3",11.5,8,4,NAVY)]:
    s=styles[name]; s.font.name="Microsoft YaHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

if PAPER_MODE:
    title = SRC.read_text(encoding="utf-8").splitlines()[0].removeprefix("# ")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(12)
    r=p.add_run(title); set_font(r,20,True,NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(16)
    r=p.add_run("综述论文稿 · 2026年8月"); set_font(r,10,False,GRAY)
else:
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(105); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("技术研究报告"); set_font(r,11,True,GOLD)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(10)
    r=p.add_run("Efficient Video Diffusion\n算法原理、负载机理与系统资源影响"); set_font(r,25,True,NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("以综述论文为主线的计算、带宽与存储影响详述"); set_font(r,13,False,BLUE)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(55)
    r=p.add_run("基于本地 paper/负载 三篇论文\n并补充截至 2026 年 8 月的公开研究"); set_font(r,11,False,GRAY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(80)
    r=p.add_run("2026 年 8 月 3 日"); set_font(r,10.5,True,NAVY)
    doc.add_page_break()

header=sec.header.paragraphs[0]; header.text="EFFICIENT VIDEO DIFFUSION  |  SURVEY MANUSCRIPT" if PAPER_MODE else "VIDEO DIFFUSION LOAD & EFFICIENCY  |  TECHNICAL REPORT"; header.alignment=WD_ALIGN_PARAGRAPH.LEFT
for r in header.runs: set_font(r,8,True,GRAY)
add_page_number(sec.footer.paragraphs[0])

lines = SRC.read_text(encoding="utf-8").splitlines()
i=0; first_title_skipped=False
while i < len(lines):
    line=lines[i].rstrip()
    if not line or line.startswith("**副标题") or line.startswith("**研究范围") or line.startswith("**版本"):
        i+=1; continue
    if line.startswith("# ") and not first_title_skipped:
        first_title_skipped=True; i+=1; continue
    if line.startswith("## "):
        p=doc.add_paragraph(style="Heading 1"); add_inline(p,line[3:]); i+=1; continue
    if line.startswith("### "):
        p=doc.add_paragraph(style="Heading 2"); add_inline(p,line[4:]); i+=1; continue
    if line.startswith("> "):
        p=doc.add_paragraph(); style_formula_paragraph(p); add_inline(p,line[2:],10.0); i+=1; continue
    if line.startswith("| "):
        rows=[]
        while i<len(lines) and lines[i].startswith("|"):
            cells=[c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?",c) for c in cells): rows.append(cells)
            i+=1
        if rows:
            table=doc.add_table(rows=len(rows), cols=len(rows[0])); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
            total=6.86; widths=[total/len(rows[0])]*len(rows[0])
            if len(widths)>=3: widths[0]=total*0.29; rem=(total-widths[0])/(len(widths)-1); widths[1:]=[rem]*(len(widths)-1)
            for ri,row in enumerate(rows):
                for ci,val in enumerate(row):
                    cell=table.cell(ri,ci); cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
                    if ri==0: shade(cell,LIGHT)
                    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
                    add_inline(p,val,9.0 if len(rows[0])>3 else 9.5)
                    for rr in p.runs:
                        if ri==0: rr.bold=True; rr.font.color.rgb=RGBColor.from_string(NAVY)
            set_repeat_table_header(table.rows[0])
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
        continue
    if re.match(r"^\d+\. ",line):
        p=doc.add_paragraph(style="List Number"); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.18); add_inline(p,re.sub(r"^\d+\. ","",line)); i+=1; continue
    if line.startswith("- "):
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.18); add_inline(p,line[2:]); i+=1; continue
    p=doc.add_paragraph(); add_inline(p,line); i+=1

core=doc.core_properties; core.title="Efficient Video Diffusion：算法原理、负载机理与系统资源影响"; core.subject="计算、带宽与存储影响论文详述版"; core.author="Codex Research"
doc.save(OUT)
print(OUT)
